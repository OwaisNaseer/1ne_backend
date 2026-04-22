"""
Text extractor implementations.
"""
import asyncio
import os
import time
from typing import Awaitable, Callable, List, Optional

from app.core.logging import get_logger
from app.domains.content_ingestion.providers.base import TextExtractor, PageText

logger = get_logger(__name__)


class PDFTextExtractor(TextExtractor):
    """
    PDF text extractor using pdfplumber.

    Runs pdfplumber in a thread-pool executor so the ASGI event loop is never
    blocked.  Reports per-page progress via an optional async callback every
    `progress_batch` pages, enabling the UI to show live "page X of Y" feedback
    for large textbooks.
    """

    provider_name = "pdfplumber"

    def __init__(self):
        try:
            import pdfplumber
            self.pdfplumber = pdfplumber
        except ImportError as e:
            logger.warning(f"pdfplumber not installed: {e}")
            self.pdfplumber = None

    def validate_config(self) -> bool:
        return self.pdfplumber is not None

    async def extract_text(
        self,
        file_path: str,
        on_progress: Optional[Callable[[int, int], Awaitable[None]]] = None,
        timeout_seconds: float = 600.0,
        progress_batch: int = 10,
        known_total_pages: Optional[int] = None,
        first_progress_timeout_seconds: Optional[float] = None,
        **kwargs,
    ) -> List[PageText]:
        """
        Extract text from a PDF file without blocking the event loop.

        Args:
            file_path:        Path to the PDF.
            on_progress:      Async callable ``(current_page, total_pages)`` called
                              every ``progress_batch`` pages.  First call always has
                              current_page=0 so the caller can record total_pages early.
            timeout_seconds:  Abort if extraction takes longer than this.
            progress_batch:   How often to call on_progress (every N pages).
            known_total_pages: If set (e.g. from fast pypdf count), emit progress before pdfplumber.open
                               so the UI is not frozen while open() blocks on huge PDFs.
            first_progress_timeout_seconds: Fail fast if no queue event arrives in this time (open stall).

        Returns:
            List of PageText objects (one per page, in order).
        """
        if not self.validate_config():
            raise RuntimeError("pdfplumber not properly configured")
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"PDF file not found: {file_path}")

        logger.info(
            "pdf_extraction_start",
            extra={"file_path": file_path, "timeout_s": timeout_seconds, "batch": progress_batch},
        )
        t0 = time.perf_counter()

        loop = asyncio.get_running_loop()
        # Thread → async bridge: the worker thread pushes tuples into this queue.
        # None sentinel signals extraction is complete (or failed).
        progress_q: asyncio.Queue[Optional[tuple[int, int]]] = asyncio.Queue()

        pdfplumber = self.pdfplumber  # local ref — avoids closure over self in thread

        def _sync_extract() -> List[PageText]:
            """Blocking pdfplumber loop; runs in the default ThreadPoolExecutor."""
            pages: List[PageText] = []
            hinted = known_total_pages
            if hinted is None:
                try:
                    from pypdf import PdfReader

                    hinted = len(PdfReader(file_path).pages)
                except Exception:
                    hinted = None
            # Critical: emit before pdfplumber.open — open() can block a long time on large PDFs while
            # the async side would otherwise show a stale coarse % with no pages_processed updates.
            if hinted is not None:
                loop.call_soon_threadsafe(progress_q.put_nowait, (0, hinted))

            try:
                with pdfplumber.open(file_path) as pdf:
                    total_pages = len(pdf.pages)
                    # If we could not pre-count, or counts disagree, sync total_pages now.
                    if hinted is None:
                        loop.call_soon_threadsafe(progress_q.put_nowait, (0, total_pages))
                    elif hinted != total_pages:
                        loop.call_soon_threadsafe(progress_q.put_nowait, (0, total_pages))

                    for page_num, page in enumerate(pdf.pages, start=1):
                        try:
                            text = page.extract_text() or ""
                            char_count = len(text.strip())
                        except Exception as page_exc:
                            logger.error(
                                "page_extraction_error",
                                extra={"file": file_path, "page": page_num, "error": str(page_exc)},
                            )
                            text, char_count = "", 0

                        pages.append(PageText(page_no=page_num, text=text, char_count=char_count))

                        # Emit progress every page when batch==1; else every N pages and at end
                        should_report = (
                            progress_batch <= 1
                            or (page_num % progress_batch == 0)
                            or page_num == total_pages
                        )
                        if should_report:
                            loop.call_soon_threadsafe(progress_q.put_nowait, (page_num, total_pages))

            finally:
                # Always signal done so the async consumer unblocks
                loop.call_soon_threadsafe(progress_q.put_nowait, None)

            return pages

        extract_task = asyncio.create_task(asyncio.to_thread(_sync_extract))

        async def _drain_and_collect() -> List[PageText]:
            got_first = False
            while True:
                if not got_first and first_progress_timeout_seconds is not None:
                    try:
                        item = await asyncio.wait_for(
                            progress_q.get(),
                            timeout=first_progress_timeout_seconds,
                        )
                    except asyncio.TimeoutError:
                        extract_task.cancel()
                        raise RuntimeError(
                            f"No PDF extraction progress within {first_progress_timeout_seconds:.0f}s. "
                            "pdfplumber may be blocked opening this file — try another PDF, split the file, "
                            "or increase EXTRACTION_FIRST_PROGRESS_TIMEOUT_SECONDS."
                        ) from None
                    got_first = True
                else:
                    item = await progress_q.get()
                    got_first = True

                if item is None:
                    break
                current_page, total_pages = item
                if on_progress:
                    try:
                        await on_progress(current_page, total_pages)
                    except Exception as cb_exc:
                        logger.warning(f"Extraction progress callback error (non-fatal): {cb_exc}")
            return await extract_task

        try:
            pages = await asyncio.wait_for(_drain_and_collect(), timeout=timeout_seconds)
        except asyncio.TimeoutError:
            extract_task.cancel()
            raise RuntimeError(
                f"PDF text extraction timed out after {timeout_seconds:.0f}s. "
                "The file may be very large, corrupted, password-protected, or have complex formatting. "
                "Try splitting the document or enabling force_ocr for image-based PDFs."
            )

        duration_ms = int((time.perf_counter() - t0) * 1000)
        logger.info(
            "pdf_extraction_complete",
            extra={"file_path": file_path, "pages": len(pages), "duration_ms": duration_ms},
        )
        return pages


class DOCXTextExtractor(TextExtractor):
    """DOCX text extractor using python-docx."""
    
    provider_name = "python-docx"
    
    def __init__(self):
        """Initialize DOCX text extractor."""
        try:
            from docx import Document
            self.Document = Document
        except ImportError as e:
            logger.warning(f"python-docx not installed: {e}")
            self.Document = None
    
    def validate_config(self) -> bool:
        """Check if python-docx is available."""
        return self.Document is not None
    
    async def extract_text(
        self,
        file_path: str,
        **kwargs
    ) -> List[PageText]:
        """
        Extract text from DOCX file.
        
        Args:
            file_path: Path to DOCX file
            
        Returns:
            List of PageText objects (DOCX doesn't have pages, so we create one page)
        """
        if not self.validate_config():
            raise RuntimeError("python-docx not properly configured")
        
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"DOCX file not found: {file_path}")
        
        logger.info(f"Extracting text from DOCX: {file_path}")
        
        try:
            doc = self.Document(file_path)
            
            # Combine all paragraphs into one text
            text_parts = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_parts.append(cell.text)
            
            full_text = "\n\n".join(text_parts)
            char_count = len(full_text.strip())
            
            # DOCX doesn't have pages, so we create one page
            pages_text = [PageText(
                page_no=1,
                text=full_text,
                char_count=char_count
            )]
            
            logger.info(f"DOCX text extraction completed: {char_count} chars")
            return pages_text
            
        except Exception as e:
            logger.error(f"DOCX text extraction failed: {e}")
            raise RuntimeError(f"Text extraction failed: {str(e)}")
