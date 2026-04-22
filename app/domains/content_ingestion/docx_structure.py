"""
Structured DOCX extraction: virtual pages + chapter_map from major headings.

Word files are common worldwide; styles vary by locale but "Heading 1" / localized
variants are widely used. When detected, we emit one PageText per section so
chunking and quiz topic strands align with document structure.
"""
from __future__ import annotations

import os
import re
from typing import Any, Dict, List, Optional, Tuple

from app.core.logging import get_logger
from app.domains.content_ingestion.providers.base import PageText
from app.domains.content_ingestion.topic_label_normalize import normalize_topic_label

logger = get_logger(__name__)


def _is_major_section_heading(style_name: str) -> bool:
    """True for top-level section breaks (Heading 1 / common localized / Title)."""
    n = (style_name or "").strip().lower()
    if not n:
        return False
    if n == "title":
        return True
    # English + variants: "Heading 1", "Heading 1 Char"
    if re.search(r"^heading\s*1(\b|_|char)", n):
        return True
    # German, Scandinavian, etc. (common Word localizations)
    if re.search(r"^überschrift\s*1(\b|_|char)", n):
        return True
    if re.search(r"^rubrik\s*1(\b|_|char)", n):
        return True
    if re.search(r"^titre\s*1(\b|_|char)", n):
        return True
    return False


def _append_tables_as_text(doc: Any, buffer: List[str]) -> None:
    """Append table cell text (order is document-global; good enough for RAG)."""
    try:
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    t = (cell.text or "").strip()
                    if t:
                        buffer.append(t)
    except Exception as e:
        logger.debug("docx_table_extract_skipped", extra={"error": str(e)})


def try_extract_docx_structured(file_path: str) -> Optional[Tuple[List[PageText], List[Dict[str, Any]]]]:
    """
    If the DOCX has at least two major sections (Heading 1 / Title breaks), return
    virtual pages + a chapter_map. Otherwise return None (caller uses legacy single-page extract).
    """
    if not file_path or not os.path.isfile(file_path):
        return None
    try:
        from docx import Document
    except ImportError:
        return None

    try:
        doc = Document(file_path)
    except Exception as e:
        logger.info("docx_structured_open_failed", extra={"path": file_path, "error": str(e)})
        return None

    sections: List[Tuple[str, List[str]]] = []
    current_title: Optional[str] = None
    buf: List[str] = []

    for p in doc.paragraphs:
        style_name = p.style.name if p.style else ""
        raw = p.text or ""
        line = raw.strip()
        if _is_major_section_heading(style_name):
            if buf:
                sections.append((current_title or "Introduction", buf[:]))
                buf = []
            current_title = normalize_topic_label(line) if line else normalize_topic_label(style_name)
            if not current_title:
                current_title = "Section"
            continue
        if line:
            buf.append(raw)

    _append_tables_as_text(doc, buf)

    if buf:
        sections.append((current_title or "Introduction", buf[:]))

    if len(sections) < 2:
        return None

    pages: List[PageText] = []
    chapter_map: List[Dict[str, Any]] = []

    for i, (title, lines) in enumerate(sections):
        body = "\n\n".join(lines).strip()
        if not body:
            continue
        pg = len(pages) + 1
        char_count = len(body)
        pages.append(PageText(page_no=pg, text=body, char_count=char_count))
        tnorm = normalize_topic_label(title, max_len=500) or f"Section {pg}"
        chapter_map.append(
            {
                "id": f"docx-sec-{pg}",
                "title": tnorm,
                "level": 1,
                "parent_id": None,
                "start_page_pdf": pg,
                "end_page_pdf": pg,
                "keywords": [],
            }
        )

    if len(pages) < 2:
        return None

    logger.info(
        "docx_structured_sections",
        extra={"path": file_path, "sections": len(pages)},
    )
    return pages, chapter_map
